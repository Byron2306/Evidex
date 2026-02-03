from __future__ import annotations

import datetime as dt
import os
from dataclasses import dataclass
from pathlib import Path

from .config import Intake
from .render import render_template


@dataclass(frozen=True)
class VendorInfo:
    name: str
    payment_instructions: str
    payment_short: str


@dataclass(frozen=True)
class InvoiceInfo:
    number: str
    date: str
    due_date: str
    currency: str
    quantity: int
    unit_price: float
    service_name: str

    @property
    def total(self) -> float:
        return float(self.quantity) * float(self.unit_price)


def load_vendor_from_env() -> VendorInfo:
    name = os.getenv("VENDOR_NAME", "EVIDEX")
    payment_link = os.getenv("PAYMENT_LINK", "").strip()
    paypal_link = os.getenv("PAYPAL_LINK", "").strip()

    # If only PayPal is set, use it as the primary link.
    primary_link = payment_link or paypal_link or "(add payment link)"

    payment_instructions = os.getenv("PAYMENT_INSTRUCTIONS", "").strip()
    if not payment_instructions:
        lines: list[str] = []
        lines.append(f"Pay via: {primary_link}")
        if paypal_link and paypal_link != primary_link:
            lines.append(f"Or pay via PayPal: {paypal_link}")
        payment_instructions = "\n".join(lines)

    payment_short = os.getenv("PAYMENT_SHORT", "").strip() or primary_link
    paypal_short = os.getenv("PAYPAL_SHORT", "").strip() or paypal_link
    if paypal_short and paypal_short != payment_short and "PayPal" not in payment_short:
        payment_short = f"{payment_short} | PayPal: {paypal_short}"
    return VendorInfo(name=name, payment_instructions=payment_instructions, payment_short=payment_short)


def _env_float(name: str, default: float) -> float:
    raw = os.getenv(name)
    if raw is None or str(raw).strip() == "":
        return float(default)
    try:
        return float(str(raw).strip())
    except ValueError:
        return float(default)


def _default_service_name() -> str:
    return os.getenv("SERVICE_NAME", "EVIDEX — Evidence Pack (48-Hour Delivery)")


def resolve_pricing_from_intake(intake: Intake) -> tuple[int, float, str, str]:
    """Return (quantity, unit_price, currency, service_name) for invoice.

    Priority:
      1) intake.billing overrides
      2) env vars by client_type
      3) fallback defaults
    """

    currency = intake.billing.currency or os.getenv("INVOICE_CURRENCY", "USD")
    service_name = intake.billing.service_name or _default_service_name()

    quantity = intake.billing.quantity if intake.billing.quantity is not None else 1

    client_type = (intake.billing.client_type or "").strip().lower()
    if client_type in {"corporate", "company", "enterprise"}:
        unit_price = _env_float("CORPORATE_UNIT_PRICE", 900.0)
        if intake.billing.unit_price is not None:
            unit_price = float(intake.billing.unit_price)
        if intake.billing.service_name is None:
            service_name = os.getenv("CORPORATE_SERVICE_NAME", service_name)
        return int(quantity), float(unit_price), str(currency), str(service_name)

    if client_type in {"consultancy", "consultant", "agency", "partner", "reseller"}:
        unit_price = _env_float("CONSULTANCY_UNIT_PRICE", 400.0)
        if intake.billing.unit_price is not None:
            unit_price = float(intake.billing.unit_price)
        if intake.billing.service_name is None:
            service_name = os.getenv("CONSULTANCY_SERVICE_NAME", service_name)
        return int(quantity), float(unit_price), str(currency), str(service_name)

    # default/ngo
    unit_price = _env_float("NGO_UNIT_PRICE", _env_float("DEFAULT_UNIT_PRICE", 500.0))
    if intake.billing.unit_price is not None:
        unit_price = float(intake.billing.unit_price)
    if intake.billing.service_name is None:
        service_name = os.getenv("NGO_SERVICE_NAME", service_name)

    return int(quantity), float(unit_price), str(currency), str(service_name)


def make_invoice_info(
    *,
    job_name: str,
    quantity: int,
    unit_price: float,
    currency: str,
    service_name: str,
    invoice_prefix: str,
) -> InvoiceInfo:
    today = dt.date.today()
    prefix = (invoice_prefix or "EVIDEX").strip() or "EVIDEX"
    number = f"{prefix}-{today.strftime('%Y%m%d')}-{job_name[:24]}".replace(" ", "-")
    return InvoiceInfo(
        number=number,
        date=today.isoformat(),
        due_date=today.isoformat(),
        currency=currency,
        quantity=quantity,
        unit_price=unit_price,
        service_name=service_name,
    )


def write_delivery_artifacts(
    *,
    intake: Intake,
    flags: list[str],
    out_dir: Path,
    job_name: str,
) -> tuple[Path, Path]:
    """Write invoice.docx and delivery email text into out_dir."""
    out_dir.mkdir(parents=True, exist_ok=True)

    templates_dir = Path(__file__).resolve().parents[2] / "templates"
    vendor = load_vendor_from_env()
    quantity, unit_price, currency, service_name = resolve_pricing_from_intake(intake)
    invoice_prefix = os.getenv("INVOICE_PREFIX", "EVIDEX")
    invoice = make_invoice_info(
        job_name=job_name,
        quantity=quantity,
        unit_price=unit_price,
        currency=currency,
        service_name=service_name,
        invoice_prefix=invoice_prefix,
    )

    invoice_md = render_template(
        templates_dir=templates_dir,
        template_rel_path="ops/invoice.md.j2",
        context={
            "vendor": vendor.__dict__,
            "client": intake.client.model_dump(),
            "pack": intake.pack.model_dump(),
            "invoice": {
                **invoice.__dict__,
                "total": invoice.total,
            },
        },
    )

    # simple docx writer (same markdown-ish convention as pack docs)
    from docx import Document

    doc = Document()
    for line in invoice_md.splitlines():
        line = line.rstrip()
        if not line:
            doc.add_paragraph("")
            continue
        if line.startswith("# "):
            doc.add_heading(line[2:], level=1)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=2)
        elif line.startswith("- "):
            doc.add_paragraph(line[2:], style="List Bullet")
        else:
            doc.add_paragraph(line)

    invoice_path = out_dir / "INVOICE.docx"
    doc.save(str(invoice_path))

    email_txt = render_template(
        templates_dir=templates_dir,
        template_rel_path="ops/delivery_email.txt.j2",
        context={
            "vendor": vendor.__dict__,
            "client": intake.client.model_dump(),
            "pack": intake.pack.model_dump(),
            "invoice": {
                **invoice.__dict__,
                "total": invoice.total,
            },
            "flags": flags,
        },
    )
    email_path = out_dir / "DELIVERY_EMAIL.txt"
    email_path.write_text(email_txt, encoding="utf-8")

    return invoice_path, email_path
