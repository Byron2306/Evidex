from __future__ import annotations

import json
import os
import shutil
import zipfile
from dataclasses import dataclass
import time
from pathlib import Path

import yaml

from .config import Intake, KPI
from .errors import TransientJobError, is_transient_file_error, with_retries
from .ingest import IngestedDoc, ingest_uploads
from .kpi_extract import extract_kpis_from_uploads
from .llm import LlmClient, load_llm_config
from .render import render_template


@dataclass(frozen=True)
class EvidenceRow:
    kpi_name: str
    target: str
    actual: str
    measurement: str
    evidence_summary: str
    sources: list[str]


def _quality_urgency_checklist(purpose: str) -> list[str]:
    p = (purpose or "").strip().lower()

    # Purpose-based, time-bound checklists to drive urgency and completeness.
    if any(k in p for k in ["tax", "vat"]):
        return [
            "Confirm tax/VAT period and filing deadline.",
            "Reconcile totals (returns ↔ source docs) and note variance explanations.",
            "Ensure invoices/receipts are readable and reference IDs match summaries.",
            "Flag missing supporting documents for any material line items.",
        ]

    if "audit" in p:
        return [
            "Map KPIs/controls to a clear audit trail (what, where, who approved).",
            "Ensure each KPI/control has at least one primary source document.",
            "Include a sampling-friendly source folder structure.",
            "List known gaps and proposed remediation actions.",
        ]

    if any(k in p for k in ["esg", "csi"]):
        return [
            "Define KPI scope + boundary (sites, subsidiaries, time period).",
            "Ensure each KPI has a calculation method + source system.",
            "Separate estimates from measured values and flag assumptions.",
            "Maintain traceable evidence for any public-facing claims.",
        ]

    if any(k in p for k in ["university", "research", "grant closeout", "closeout"]):
        return [
            "Confirm sponsor closeout requirements and submission checklist.",
            "Verify KPI/outputs align with approved proposal and reporting template.",
            "Ensure supporting documents include approvals, ethics letters, or protocols if applicable.",
            "List any deviations and attach justification/approvals.",
        ]

    if any(k in p for k in ["donor", "grant", "funder"]):
        return [
            "Match KPIs exactly to the donor logframe/template naming.",
            "Ensure each KPI has evidence + a conservative narrative justification.",
            "Flag missing months/documents early (so client can backfill).",
            "Include finance summaries that tie to receipts/invoices where relevant.",
        ]

    # Default
    return [
        "Confirm reporting period + deadline.",
        "Ensure each KPI has at least one supporting source.",
        "List known gaps and what to provide next.",
    ]


def _build_quality_report(*, intake: Intake, evidence_rows: list[EvidenceRow], docs: list[IngestedDoc], flags: list[str]) -> str:
    total_kpis = len(evidence_rows)
    kpis_with_sources = sum(1 for r in evidence_rows if r.sources)
    coverage_pct = (100.0 * kpis_with_sources / total_kpis) if total_kpis else 0.0
    missing_kpis = [r.kpi_name for r in evidence_rows if not r.sources]

    # Heuristic: treat any "Missing/weak evidence" flags as evidence gaps.
    evidence_gap_flags = [f for f in flags if f.lower().startswith("missing/weak evidence")]

    lines: list[str] = []
    lines.append("QUALITY REPORT")
    lines.append("============")
    lines.append("")
    lines.append(f"Organization: {intake.client.organization}")
    lines.append(f"Project: {intake.pack.project_name}")
    lines.append(f"Purpose: {intake.pack.purpose}")
    lines.append(f"Period: {intake.pack.reporting_period.start} to {intake.pack.reporting_period.end}")
    lines.append("")

    lines.append("Coverage")
    lines.append("--------")
    lines.append(f"- Uploaded source files parsed: {len([d for d in docs if d.path.exists()])}")
    lines.append(f"- KPIs in pack: {total_kpis}")
    lines.append(f"- KPIs with matched evidence sources: {kpis_with_sources} ({coverage_pct:.0f}%)")
    if missing_kpis:
        lines.append("- KPIs missing matched sources:")
        for name in missing_kpis[:25]:
            lines.append(f"  - {name}")
        if len(missing_kpis) > 25:
            lines.append(f"  - (+{len(missing_kpis) - 25} more)")
    lines.append("")

    lines.append("Flags")
    lines.append("-----")
    if flags:
        for f in flags:
            lines.append(f"- {f}")
    else:
        lines.append("- No material issues flagged.")
    lines.append("")

    lines.append("Urgency checklist")
    lines.append("-----------------")
    for item in _quality_urgency_checklist(intake.pack.purpose):
        lines.append(f"- {item}")
    lines.append("")

    lines.append("Next actions")
    lines.append("-----------")
    if evidence_gap_flags or missing_kpis:
        lines.append("- Provide missing documents for the KPIs listed above (or adjust KPI wording to match your source docs).")
        lines.append("- If you have a KPI spreadsheet/logframe, upload it to improve matching.")
    lines.append("- Confirm the pack purpose and any required template headings (donor/auditor/board).")
    lines.append("- Keep this report with the ZIP as an internal QA record.")

    return "\n".join(lines) + "\n"


def load_intake(path: Path) -> Intake:
    raw_text = path.read_text(encoding="utf-8")
    if path.suffix.lower() in {".yaml", ".yml"}:
        raw = yaml.safe_load(raw_text)
    else:
        raw = json.loads(raw_text)
    return Intake.model_validate(raw)


def _score_match(kpi: KPI, doc: IngestedDoc) -> float:
    if not doc.text:
        return 0.0
    needle = kpi.name.lower()
    hay = doc.text.lower()
    if needle in hay:
        return 10.0
    # simple keyword overlap
    terms = [t for t in needle.replace("/", " ").replace("-", " ").split() if len(t) > 3]
    if not terms:
        return 0.0
    hits = sum(1 for t in terms if t in hay)
    return hits / max(1, len(terms))


def build_evidence_rows(intake: Intake, docs: list[IngestedDoc], llm: LlmClient) -> tuple[list[EvidenceRow], list[str]]:
    rows: list[EvidenceRow] = []
    flags: list[str] = []

    summary_use_llm = str(os.getenv("SUMMARY_USE_LLM", "1")).strip().lower() in {"1", "true", "yes", "y", "on"}

    for kpi in intake.kpis:
        scored = sorted((( _score_match(kpi, d), d) for d in docs), key=lambda x: x[0], reverse=True)
        top = [d for s, d in scored if s > 0][:3]
        source_names = [d.path.name for d in top] if top else []

        if summary_use_llm and llm.enabled and docs:
            summary_model = os.getenv("LLM_MODEL_SUMMARY") or os.getenv("OPENAI_MODEL_SUMMARY") or os.getenv("OLLAMA_MODEL_SUMMARY")
            system = "You generate concise audit-ready evidence summaries."
            user = (
                "KPI: " + kpi.name + "\n"
                "Target: " + str(kpi.target) + "\n"
                "Actual: " + str(kpi.actual) + "\n"
                "Measurement: " + kpi.measurement + "\n\n"
                "Summarize the best available evidence for this KPI in 2-4 sentences. "
                "Be conservative; do not invent facts. If evidence is missing, say so explicitly."\
                "\n\n"
                "Available sources (filename + excerpt):\n" + "\n\n".join(
                    f"- {d.path.name}:\n{(d.text or '')[:1200]}" for d in top
                )
            )
            summary = llm.complete(system=system, user=user, model=summary_model)
            if not summary:
                if top:
                    summary = f"Evidence found in {', '.join(source_names)} supporting measurement: {kpi.measurement}."
                else:
                    summary = "No matching evidence found in uploaded sources; confirm missing documents or adjust KPI wording."
                    flags.append(f"Missing/weak evidence for KPI: {kpi.name}")
        else:
            if top:
                summary = f"Evidence found in {', '.join(source_names)} supporting measurement: {kpi.measurement}."
            else:
                summary = "No matching evidence found in uploaded sources; confirm missing documents or adjust KPI wording."
                flags.append(f"Missing/weak evidence for KPI: {kpi.name}")

        rows.append(
            EvidenceRow(
                kpi_name=kpi.name,
                target=str(kpi.target),
                actual=str(kpi.actual),
                measurement=kpi.measurement,
                evidence_summary=summary.strip(),
                sources=source_names,
            )
        )

    # add known gaps as flags
    for gap in intake.constraints.known_gaps:
        flags.append(f"Known gap: {gap}")

    return rows, flags


def _write_docx(path: Path, markdown_like_text: str) -> None:
    from docx import Document

    doc = Document()
    for line in markdown_like_text.splitlines():
        line = line.rstrip()
        if not line:
            doc.add_paragraph("")
            continue
        if line.startswith("# "):
            doc.add_heading(line[2:], level=1)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=2)
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=3)
        elif line.startswith("- "):
            doc.add_paragraph(line[2:], style="List Bullet")
        else:
            doc.add_paragraph(line)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))


def write_pack(*, intake_path: Path, uploads_dir: Path, out_dir: Path, job_name: str | None = None) -> tuple[Path, Path, list[str]]:
    intake = load_intake(intake_path)

    # If the client uploaded a KPI sheet (preferred) but did not paste KPIs,
    # auto-extract KPI rows from CSV/XLSX so the evidence table isn't empty.
    if not intake.kpis:
        extracted_kpis, _provenance = extract_kpis_from_uploads(uploads_dir)
        if extracted_kpis:
            intake = intake.model_copy(update={"kpis": extracted_kpis})

    docs = ingest_uploads(uploads_dir)
    llm = LlmClient(load_llm_config())

    templates_dir = Path(__file__).resolve().parents[2] / "templates"

    evidence_rows, flags = build_evidence_rows(intake, docs, llm)

    def _slugify(s: str) -> str:
        s2 = "".join(c if c.isalnum() or c in {"-", "_"} else "_" for c in (s or ""))
        s2 = "_".join(p for p in s2.split("_") if p)  # collapse repeats
        return s2[:140]

    org = (intake.client.organization or "").strip()
    proj = (intake.pack.project_name or intake.pack.purpose or "").strip()
    start = (intake.pack.reporting_period.start or "").strip()
    end = (intake.pack.reporting_period.end or "").strip()

    # If org/project are blank (common when intake extraction failed), the naive slug becomes "to".
    # In that case, fall back to job folder name to avoid collisions and locked-folder retries.
    base_slug = _slugify(f"{org}__{proj}__{start}_to_{end}") if (org or proj) else ""
    job_slug = _slugify((job_name or "").strip())

    # If intake fields are missing (common when a form title mismatch happens),
    # fall back to the job folder name so we don't create collisions like "_____to_".
    safe_slug = base_slug or job_slug or _slugify(intake_path.parent.name) or f"job_{int(__import__('time').time())}"

    pack_root = out_dir / safe_slug

    def _rm_tree(target: Path) -> None:
        if target.exists():
            shutil.rmtree(target)

    try:
        with_retries(lambda: _rm_tree(pack_root), attempts=8)
    except TransientJobError:
        # If sync/AV/Explorer has a handle open and we can't cleanly delete, don't strand the job.
        # Instead, write to a new unique folder.
        pack_root = out_dir / f"{safe_slug}__{int(time.time())}"
    except OSError as e:
        if is_transient_file_error(e):
            # Same idea as above, but handle the (rare) case where retry wrapper didn't classify it.
            pack_root = out_dir / f"{safe_slug}__{int(time.time())}"
        else:
            raise

    # folder structure
    (pack_root / "00_EXEC_SUMMARY").mkdir(parents=True, exist_ok=True)
    (pack_root / "01_EVIDENCE_TABLE").mkdir(parents=True, exist_ok=True)
    (pack_root / "02_NARRATIVE").mkdir(parents=True, exist_ok=True)
    (pack_root / "03_SOURCES").mkdir(parents=True, exist_ok=True)

    if intake.pack.include_appendix:
        (pack_root / "04_APPENDIX").mkdir(parents=True, exist_ok=True)

    # executive summary
    exec_md = render_template(
        templates_dir=templates_dir,
        template_rel_path="grant_report/executive_summary.md.j2",
        context={
            "client": intake.client.model_dump(),
            "pack": intake.pack.model_dump(),
            "kpis": [k.model_dump() for k in intake.kpis],
            "constraints": intake.constraints.model_dump(),
            "source_count": len([d for d in docs if d.path.exists()]),
        },
    )
    _write_docx(pack_root / "00_EXEC_SUMMARY" / "executive_summary.docx", exec_md)

    # evidence table
    import pandas as pd

    df = pd.DataFrame(
        [
            {
                "KPI": r.kpi_name,
                "Target": r.target,
                "Actual": r.actual,
                "Measurement": r.measurement,
                "Evidence summary": r.evidence_summary,
                "Sources": "; ".join(r.sources),
            }
            for r in evidence_rows
        ]
    )
    df.to_csv(pack_root / "01_EVIDENCE_TABLE" / "evidence_table.csv", index=False)
    df.to_excel(pack_root / "01_EVIDENCE_TABLE" / "evidence_table.xlsx", index=False)

    # narrative
    narrative_use_llm = str(os.getenv("NARRATIVE_USE_LLM", "0")).strip().lower() in {"1", "true", "yes", "y", "on"}
    narrative_text = ""
    if narrative_use_llm and llm.enabled and evidence_rows:
        narrative_model = os.getenv("LLM_MODEL_NARRATIVE") or os.getenv("OPENAI_MODEL_NARRATIVE") or os.getenv("OLLAMA_MODEL_NARRATIVE")
        system = (
            "You write conservative, audit-ready narrative justifications for evidence packs. "
            "Do not invent facts. If evidence is weak or missing, explicitly say so. "
            "Use a professional tone suitable for auditors/finance/compliance."
        )
        user = (
            f"Purpose: {intake.pack.purpose}\n"
            f"Project: {intake.pack.project_name}\n"
            f"Reporting period: {intake.pack.reporting_period.start} to {intake.pack.reporting_period.end}\n\n"
            "Write a short narrative (max ~1.5 pages) with these sections:\n"
            "1) Overview\n2) Results against KPIs (bullet per KPI)\n3) Data integrity & audit trail\n4) Risks/flags\n\n"
            "KPI rows (name, target, actual, measurement, evidence summary, sources):\n"
            + "\n".join(
                f"- KPI: {r.kpi_name} | Target: {r.target} | Actual: {r.actual} | Measurement: {r.measurement} | "
                f"Evidence: {r.evidence_summary} | Sources: {', '.join(r.sources) if r.sources else 'None'}"
                for r in evidence_rows
            )
            + "\n\n"
            + ("Flags:\n" + "\n".join(f"- {f}" for f in flags) if flags else "Flags:\n- None")
        )
        narrative_text = llm.complete(system=system, user=user, model=narrative_model)

    if not narrative_text:
        narrative_text = render_template(
            templates_dir=templates_dir,
            template_rel_path="grant_report/narrative.md.j2",
            context={
                "pack": intake.pack.model_dump(),
                "evidence_rows": [r.__dict__ for r in evidence_rows],
                "flags": flags,
            },
        )
    _write_docx(pack_root / "02_NARRATIVE" / "narrative_justification.docx", narrative_text)

    # sources copy
    for d in docs:
        try:
            rel = d.path.relative_to(uploads_dir)
        except Exception:
            rel = Path(d.path.name)
        dest = pack_root / "03_SOURCES" / rel
        dest.parent.mkdir(parents=True, exist_ok=True)
        def _do_copy(src: Path = d.path, dst: Path = dest) -> None:
            shutil.copy2(src, dst)

        try:
            with_retries(_do_copy, attempts=8)
        except TransientJobError as e:
            raise TransientJobError(f"File appears locked during copy: {d.path} -> {dest}. {e}")
        except OSError as e:
            # Preserve a clearer message for Windows sharing violations if retries weren't used.
            if is_transient_file_error(e):
                raise TransientJobError(f"File appears locked during copy: {d.path} -> {dest}. {e}")
            raise

    # appendix (optional)
    if intake.pack.include_appendix:
        appendix_path = pack_root / "04_APPENDIX" / "appendix_source_summaries.docx"
        if llm.enabled and docs:
            system = "You summarize documents for an evidence pack appendix."
            user = "Summarize each file in 3-6 bullets. Be factual; do not invent.\n\n" + "\n\n".join(
                f"FILE: {d.path.name}\nEXCERPT:\n{(d.text or '')[:1800]}" for d in docs[:12]
            )
            appendix_text = llm.complete(system=system, user=user)
            if not appendix_text:
                appendix_text = "No appendix summaries generated."
        else:
            appendix_text = "Appendix summaries disabled (no LLM configured) or no parsable text sources."\

        _write_docx(appendix_path, "# Appendix — Source Summaries\n\n" + appendix_text)

    # delivery note
    delivery = (
        "Evidence Pack Delivery\n"
        "====================\n\n"
        "This folder is structured to be audit-ready:\n"
        "- 00_EXEC_SUMMARY: executive overview\n"
        "- 01_EVIDENCE_TABLE: KPI → evidence → sources\n"
        "- 02_NARRATIVE: narrative justification\n"
        "- 03_SOURCES: copied source files (audit trail)\n"
        "- 04_APPENDIX: optional summaries\n"
    )
    (pack_root / "README_DELIVERY.txt").write_text(delivery, encoding="utf-8")

    # quality report (deterministic; no LLM required)
    quality_text = _build_quality_report(intake=intake, evidence_rows=evidence_rows, docs=docs, flags=flags)
    (pack_root / "QUALITY_REPORT.txt").write_text(quality_text, encoding="utf-8")

    # invoice + delivery email (include in ZIP by generating before zipping)
    try:
        from .ops import write_delivery_artifacts

        resolved_job_name = (job_name or intake_path.stem or "job").strip() or "job"
        write_delivery_artifacts(intake=intake, flags=flags, out_dir=pack_root, job_name=resolved_job_name)
    except Exception:
        # Non-fatal: pack can still be delivered without invoice artifacts.
        pass

    # zip
    # Avoid collisions when multiple jobs share the same intake-derived slug (or when retries create
    # multiple job folders). The job folder name is stable and unique for downstream syncing.
    zip_slug = safe_slug
    if job_slug and job_slug not in zip_slug:
        zip_slug = _slugify(f"{safe_slug}__{job_slug}")
    zip_path = out_dir / f"{zip_slug}.zip"
    tmp_zip_path = out_dir / f"{zip_slug}.zip.tmp"

    def _safe_unlink(p: Path) -> None:
        if p.exists():
            p.unlink()

    try:
        with_retries(lambda: _safe_unlink(tmp_zip_path), attempts=6)
        with_retries(lambda: _safe_unlink(zip_path), attempts=6)
    except TransientJobError as e:
        raise TransientJobError(f"Zip output path appears locked: {zip_path}. {e}")

    def _write_tmp_zip() -> None:
        with zipfile.ZipFile(tmp_zip_path, "w", compression=zipfile.ZIP_DEFLATED) as z:
            for p in pack_root.rglob("*"):
                if p.is_file():
                    z.write(p, arcname=str(p.relative_to(pack_root)))

    try:
        with_retries(_write_tmp_zip, attempts=6)
        with_retries(lambda: tmp_zip_path.replace(zip_path), attempts=6)
    except TransientJobError as e:
        raise TransientJobError(f"Failed to write/rename zip (likely locked by sync): {zip_path}. {e}")

    return zip_path, pack_root, flags
